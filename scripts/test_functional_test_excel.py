#!/usr/bin/env python3
"""Synthetic tests for the user-supplied fixed Excel functional-test layout."""
import json
import copy
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from datetime import date
from unittest.mock import patch
from xml.etree import ElementTree as ET
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills/feature-map-workflow/scripts"))
from human_work_plan import normalize_v2, case_checkpoints
SCRIPT = ROOT / "scripts" / "build_functional_test_excel.py"

class WorkbookTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.template = self.root / "template.xlsx"
        w = Workbook()
        cover = w.active
        cover.title = "表紙"
        listing = w.create_sheet("検証リスト")
        detail = w.create_sheet("FT-01")
        detail["H30"] = '=IF(F6="対象","未実施","未判定")'
        detail["B31"] = '=IF(B16="","未設計","設計済")'
        detail["B5"] = "FT-01"
        detail["A40"] = "E-01"
        listing["A7"] = "='FT-01'!B5"
        w.save(self.template)
        self.xml = self.root / "feature-map.xml"
        self.xml.write_text("""<featureMap version="1.3" mode="change" state="draft">
<meta><system>架空システム</system><feature>受付</feature></meta>
<goal><purpose>依頼受付</purpose></goal>
<useCase id="UC-ONE"><scenario>依頼する</scenario><input>商品ID</input></useCase>
<verify><case id="V-ONE" type="normal" status="not-run">
<condition>通常入力</condition><expect>受付番号を表示</expect>
</case></verify></featureMap>""", encoding="utf-8")

    def cli(self, *args):
        return subprocess.run([sys.executable, str(SCRIPT), *map(str,args)],
                              capture_output=True,text=True,timeout=25)

    def plan(self):
        target = self.root / "plan.xlsx"
        proc = self.cli("plan","--template",self.template,"--feature-map",self.xml,
                        "--revision","sha-test","--output",target)
        self.assertEqual(0,proc.returncode,proc.stderr)
        return target

    def test_plan_preserves_formula_and_pending_decision(self):
        target=self.plan()
        w=load_workbook(target)
        self.assertEqual("f",w["FT-01"]["H30"].data_type)
        self.assertEqual("f",w["FT-01"]["B31"].data_type)
        self.assertEqual("未判定",w["FT-01"]["F6"].value)
        self.assertEqual("通常入力",w["FT-01"]["B16"].value)
        self.assertIsNone(w["FT-01"]["H16"].value)
        self.assertIn("VERIFY:V-ONE",w["FT-01"]["B7"].value)
        self.assertEqual("sha-test",w["表紙"]["K19"].value)

    def test_case_count_above_ten_preserves_footer(self):
        tree=ET.parse(self.xml)
        node=tree.getroot().find("verify")
        node.clear()
        for idx in range(12):
            c=ET.SubElement(node,"case",id=f"V-{idx}",type="normal",status="not-run")
            ET.SubElement(c,"condition").text=f"case {idx}"
            ET.SubElement(c,"expect").text="returns OK"
        tree.write(self.xml,encoding="utf-8")
        plan=self.plan()
        w=load_workbook(plan)
        self.assertIn("FT-21",w.sheetnames)
        self.assertEqual("='FT-21'!B5",w["検証リスト"]["A27"].value)

    def test_result_requires_matching_revision_and_step_evidence(self):
        target=self.plan()
        p=self.root/"result.json"
        base=dict(schemaVersion=1,runId="R-1",verifyCaseId="V-ONE",
                  sourceRevision="different",status="passed",checks=[],evidence=[])
        p.write_text(json.dumps(base),encoding="utf-8")
        output=self.root/"resulted.xlsx"
        bad=self.cli("results","--workbook",target,"--result",p,"--output",output)
        self.assertNotEqual(0,bad.returncode)
        self.assertFalse(output.exists())
        base["sourceRevision"]="sha-test"
        p.write_text(json.dumps(base),encoding="utf-8")
        good=self.cli("results","--workbook",target,"--result",p,"--output",output)
        self.assertEqual(0,good.returncode,good.stderr)
        w=load_workbook(output)
        self.assertIsNone(w["FT-01"]["H16"].value)
        self.assertIn("R-1",w["FT-01"]["B26"].value)
        base["checks"]=[dict(name="legacy observation",ok=True)]
        p.write_text(json.dumps(base),encoding="utf-8")
        legacy=self.root/"legacy-unmapped.xlsx"
        self.assertEqual(0,self.cli("results","--workbook",target,"--result",p,"--output",legacy).returncode)
        self.assertIsNone(load_workbook(legacy)["FT-01"]["H16"].value)

    def test_formula_injection_and_dtd_are_refused(self):
        tree=ET.parse(self.xml)
        tree.getroot().find("./verify/case/condition").text="=HYPERLINK(1)"
        tree.write(self.xml,encoding="utf-8")
        w=load_workbook(self.plan())
        self.assertTrue(w["FT-01"]["B16"].value.startswith("'="))
        self.xml.write_text("<!DOCTYPE featureMap><featureMap/>",encoding="utf-8")
        failed=self.cli("plan","--template",self.template,"--feature-map",self.xml,
                        "--revision","sha-test","--output",self.root/"other.xlsx")
        self.assertNotEqual(0,failed.returncode)

    def v2_plan(self, name="v2.xlsx", mutate=None, environment=None):
        tree=ET.parse(self.xml)
        if tree.getroot().find("./verify/case[@id='V-TWO']") is None:
            c=ET.SubElement(tree.getroot().find("verify"),"case",id="V-TWO",type="error",status="not-run")
            ET.SubElement(c,"condition").text="Invalid input"
            ET.SubElement(c,"expect").text="Error without registration"
        tree.write(self.xml,encoding="utf-8")
        plan=json.loads((ROOT/"examples/fictional-stock-request/human-work-plan.v2.example.json").read_text(encoding="utf-8"))
        plan["sourceRevision"]="sha-test"
        plan["requiredConditions"]={"V-ONE":plan["requiredConditions"]["V-SUBMIT"],
                                    "V-TWO":plan["requiredConditions"]["V-REJECT"]}
        for part in plan["procedures"]:
            part["appliesTo"]["sourceRevision"]="sha-test"
        for session in plan["sessions"]:
            for step in session["steps"]:
                for check in step["checks"]:
                    check["caseId"]={"V-SUBMIT":"V-ONE","V-REJECT":"V-TWO"}[check["caseId"]]
        if mutate:mutate(plan)
        source=self.root/"human-plan.json"
        source.write_text(json.dumps(plan),encoding="utf-8")
        target=self.root/name
        extra=["--environment",environment] if environment is not None else []
        run=self.cli("plan","--template",self.template,"--feature-map",self.xml,
                     "--revision","sha-test","--human-plan",source,"--output",target,*extra)
        return target,run

    def test_v2_environment_matches_procedure_applicability_and_cover(self):
        target,run=self.v2_plan()
        self.assertEqual(0,run.returncode,run.stderr)
        self.assertEqual("fictional-isolated",load_workbook(target)["表紙"]["D19"].value)
        target,run=self.v2_plan(name="incompatible.xlsx",environment="other-environment")
        self.assertNotEqual(0,run.returncode)
        self.assertFalse(target.exists())

    def test_v2_shared_operations_project_to_independent_pending_case_rows(self):
        target,run=self.v2_plan()
        self.assertEqual(0,run.returncode,run.stderr)
        w=load_workbook(target)
        self.assertIn("CONDITION:NORMAL-READY",w["FT-01"]["B16"].value)
        self.assertIn("CONDITION:NORMAL-RESULT",w["FT-01"]["B17"].value)
        self.assertIn("CONDITION:ERROR-READY",w["FT-07"]["B16"].value)
        self.assertIn("OPERATION:OP-OPEN",w["FT-01"]["B16"].value)
        self.assertIn("OPERATION:OP-OPEN",w["FT-07"]["B16"].value)
        for sheet in ("FT-01","FT-07"):
            self.assertEqual("未判定",w[sheet]["F6"].value)
            self.assertEqual("f",w[sheet]["H30"].data_type)
            self.assertTrue(all(w[sheet][f"H{row}"].value is None for row in (16,17)))

    def test_v2_rejects_revision_or_row_overflow_before_creating_workbook(self):
        def overflow(plan):
            operation=plan["sessions"][0]["steps"][-1]
            original=operation["checks"][0]
            for number in range(3,8):
                check=dict(original,stepNo=number,conditionId=f"EXTRA-{number}")
                operation["checks"].append(check)
                plan["requiredConditions"][check["caseId"]].append(check["conditionId"])
        for index,mutate in enumerate((
            lambda p:p.update(sourceRevision="different"),
            overflow,
        )):
            target,run=self.v2_plan(name=f"bad-{index}.xlsx",mutate=mutate)
            self.assertNotEqual(0,run.returncode)
            self.assertFalse(target.exists())

    def test_v2_result_keeps_case_status_and_refuses_wrong_condition_or_missing_evidence(self):
        target,run=self.v2_plan()
        self.assertEqual(0,run.returncode,run.stderr)
        for final,label in (("failed","不合格"),("blocked","保留")):
            (self.root/"ready.txt").write_text("ready observation",encoding="utf-8")
            (self.root/"result.txt").write_text("result observation",encoding="utf-8")
            result=dict(schemaVersion=1,runId="RUN-1",verifyCaseId="V-ONE",sourceRevision="sha-test",
                        status=final,checks=[
                            dict(stepNo=1,conditionId="NORMAL-READY",status="passed",evidence=["ready.txt"]),
                            dict(stepNo=2,conditionId="NORMAL-RESULT",status=final,evidence=["result.txt"])],
                        evidence=["ready.txt","result.txt"])
            source=self.root/"result.json";source.write_text(json.dumps(result),encoding="utf-8")
            output=self.root/f"{final}.xlsx"
            imported=self.cli("results","--workbook",target,"--result",source,"--output",output)
            self.assertEqual(0,imported.returncode,imported.stderr)
            w=load_workbook(output)
            self.assertEqual("合格",w["FT-01"]["H16"].value)
            self.assertEqual(label,w["FT-01"]["H17"].value)
            self.assertIsNone(w["FT-07"]["H16"].value)
            for key,value in (("conditionId","WRONG"),("evidence",["missing.txt"])):
                result["checks"][1][key]=value
                source.write_text(json.dumps(result),encoding="utf-8")
                invalid=self.root/f"bad-{final}-{key}.xlsx"
                self.assertNotEqual(0,self.cli("results","--workbook",target,"--result",source,"--output",invalid).returncode)
                self.assertFalse(invalid.exists())
                result["checks"][1].update(conditionId="NORMAL-RESULT",evidence=["result.txt"])
        result.update(status="not-run",checks=[],evidence=[])
        source.write_text(json.dumps(result),encoding="utf-8")
        notrun=self.root/"not-run.xlsx"
        self.assertEqual(0,self.cli("results","--workbook",target,"--result",source,"--output",notrun).returncode)
        self.assertIsNone(load_workbook(notrun)["FT-01"]["H16"].value)

    def test_v2_pass_requires_all_planned_checkpoints(self):
        target,run=self.v2_plan()
        self.assertEqual(0,run.returncode,run.stderr)
        for name in ("ready.txt","result.txt"):
            (self.root/name).write_text("synthetic observation",encoding="utf-8")
        checks=[dict(stepNo=1,conditionId="NORMAL-READY",status="passed",evidence=["ready.txt"]),
                dict(stepNo=2,conditionId="NORMAL-RESULT",status="passed",evidence=["result.txt"])]
        for count in (0,1,2):
            with self.subTest(count=count):
                record=dict(schemaVersion=1,runId="R-PASS",verifyCaseId="V-ONE",sourceRevision="sha-test",
                            status="passed",checks=checks[:count],evidence=["ready.txt","result.txt"])
                source=self.root/f"pass-{count}.json";source.write_text(json.dumps(record),encoding="utf-8")
                output=self.root/f"pass-{count}.xlsx"
                result=self.cli("results","--workbook",target,"--result",source,"--output",output)
                if count<2:
                    self.assertNotEqual(0,result.returncode)
                    self.assertIn("every planned v2 checkpoint",result.stderr)
                    self.assertFalse(output.exists())
                else:
                    self.assertEqual(0,result.returncode,result.stderr)
                    ws=load_workbook(output)["FT-01"]
                    self.assertEqual(["合格","合格"],[ws["H16"].value,ws["H17"].value])
        record["checks"]=checks+[dict(status="failed",name="unmapped assertion")]
        source=self.root/"contradictory.json";source.write_text(json.dumps(record),encoding="utf-8")
        output=self.root/"contradictory.xlsx"
        self.assertNotEqual(0,self.cli("results","--workbook",target,"--result",source,"--output",output).returncode)
        self.assertFalse(output.exists())

    def test_v1_free_text_prefix_is_not_a_v2_checkpoint(self):
        for index,condition in enumerate(("CONDITION: user is authorized",
                                           "CONDITION:X | OPERATION: described behavior")):
            with self.subTest(condition=condition):
                tree=ET.parse(self.xml);tree.getroot().find("./verify/case/condition").text=condition
                tree.write(self.xml,encoding="utf-8")
                target=self.root/f"legacy-plan-{index}.xlsx"
                planned=self.cli("plan","--template",self.template,"--feature-map",self.xml,
                                 "--revision","sha-test","--output",target)
                self.assertEqual(0,planned.returncode,planned.stderr)
                source=self.root/f"legacy-result-{index}.json"
                source.write_text(json.dumps(dict(schemaVersion=1,runId="LEGACY",verifyCaseId="V-ONE",
                                                  sourceRevision="sha-test",status="passed",checks=[],evidence=[])),encoding="utf-8")
                output=self.root/f"legacy-import-{index}.xlsx"
                result=self.cli("results","--workbook",target,"--result",source,"--output",output)
                self.assertEqual(0,result.returncode,result.stderr)
                self.assertIsNone(load_workbook(output)["FT-01"]["H16"].value)

    def test_v2_keeps_authoritative_expectation_next_to_checkpoint(self):
        def contradict(plan):
            for session in plan["sessions"]:
                for operation in session["steps"]:
                    for check in operation["checks"]:
                        if check["caseId"]=="V-TWO":check["expected"]="INVALID REQUEST IS ACCEPTED"
        target,run=self.v2_plan(name="oracle.xlsx",mutate=contradict)
        self.assertEqual(0,run.returncode,run.stderr)
        ws=load_workbook(target)["FT-07"]
        expected=" ".join(str(ws[f"{col}{row}"].value or "") for row in (16,17) for col in ("D","E"))
        self.assertIn("Error without registration",expected)
        self.assertIn("INVALID REQUEST IS ACCEPTED",expected)
        self.assertIn("ケース全体の期待結果",expected)
        self.assertEqual("未判定",ws["F6"].value)

    def test_v2_case_id_normalization_is_shared_with_projections(self):
        def whitespace(plan):
            for session in plan["sessions"]:
                for operation in session["steps"]:
                    for check in operation["checks"]:
                        if check["caseId"]=="V-ONE":check["caseId"]=" V-ONE "
        target,run=self.v2_plan(name="normalized.xlsx",mutate=whitespace)
        self.assertEqual(0,run.returncode,run.stderr)
        raw=json.loads((self.root/"human-plan.json").read_text(encoding="utf-8"))
        normalized=normalize_v2(raw,["V-ONE","V-TWO"],"sha-test")
        self.assertEqual({"V-ONE","V-TWO"},set(case_checkpoints(normalized)))
        self.assertTrue(any(c["caseId"]==" V-ONE " for s in raw["sessions"] for op in s["steps"] for c in op["checks"]))
        self.assertTrue(all(c["caseId"] in {"V-ONE","V-TWO"} for s in normalized["sessions"] for op in s["steps"] for c in op["checks"]))

    def test_v2_reuse_rejects_future_verification_date(self):
        self.v2_plan()
        raw=json.loads((self.root/"human-plan.json").read_text(encoding="utf-8"))
        raw["sessions"][0]["steps"][0]["trial"]=dict(mode="reuse",reason="dated trial",recordRef="fictional:trial")
        class FixedDate(date):
            @classmethod
            def today(cls):return cls(2026,10,24)
        with patch("human_work_plan.date",FixedDate):
            raw["procedures"][0]["verifiedAt"]="2026-10-24"
            normalize_v2(raw,["V-ONE","V-TWO"],"sha-test")
            raw["procedures"][0]["verifiedAt"]="2026-10-25"
            with self.assertRaisesRegex(ValueError,"future"):
                normalize_v2(raw,["V-ONE","V-TWO"],"sha-test")

    def test_v2_excel_rejects_cross_session_case_and_keeps_roles(self):
        def split(plan):
            first=plan["sessions"][0];second=copy.deepcopy(first)
            second["id"]="SECOND";second["actor"]="operator2"
            second["steps"]=first["steps"][1:];first["steps"]=first["steps"][:1]
            first.update(sharedSetup=["FIRST SETUP"],stopOn=["FIRST STOP"],cleanup=["FIRST CLEANUP"])
            second.update(sharedSetup=["SECOND SETUP"],stopOn=["SECOND STOP"],cleanup=["SECOND CLEANUP"])
            for operation in second["steps"]:
                operation["dependsOn"]=[];operation["restartFrom"]=operation["id"]
            for part in plan["procedures"][1:]:part["appliesTo"]["actor"]="operator2"
            plan["sessions"]=[first,second]
        target,run=self.v2_plan(name="cross-session.xlsx",mutate=split)
        self.assertNotEqual(0,run.returncode)
        self.assertIn("across sessions",run.stderr)
        self.assertFalse(target.exists())
        raw=json.loads((self.root/"human-plan.json").read_text(encoding="utf-8"))
        self.assertEqual(2,len(normalize_v2(raw,["V-ONE","V-TWO"],"sha-test")["sessions"]))
        target,run=self.v2_plan(name="roles.xlsx")
        self.assertEqual(0,run.returncode,run.stderr)
        raw=json.loads((self.root/"human-plan.json").read_text(encoding="utf-8"));session=raw["sessions"][0]
        context=load_workbook(target)["FT-01"]["B8"].value
        for value in [session["id"],session["actor"],*session["stopOn"],*session["cleanup"]]:
            self.assertIn(value,context)
        def separate_cases(plan):
            first=plan["sessions"][0];second=copy.deepcopy(first)
            first["steps"]=first["steps"][:2]
            first["steps"][0]["checks"]=[c for c in first["steps"][0]["checks"] if c["caseId"]=="V-ONE"]
            second["id"]="SECOND"
            second["steps"]=[second["steps"][0],second["steps"][2]]
            opening=second["steps"][0];opening["id"]="OP-OPEN-SECOND";opening["restartFrom"]=opening["id"]
            opening["checks"]=[c for c in opening["checks"] if c["caseId"]=="V-TWO"]
            second["steps"][1]["dependsOn"]=[opening["id"]]
            second["steps"][1]["restartFrom"]=opening["id"]
            plan["sessions"]=[first,second]
        target,run=self.v2_plan(name="separate-sessions.xlsx",mutate=separate_cases)
        self.assertEqual(0,run.returncode,run.stderr)
        book=load_workbook(target)
        self.assertNotIn("SESSION:SECOND",book["FT-01"]["B8"].value)
        self.assertIn("SESSION:SECOND",book["FT-07"]["B8"].value)

if __name__=="__main__":
    unittest.main()
