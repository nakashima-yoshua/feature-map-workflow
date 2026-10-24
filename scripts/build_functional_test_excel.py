#!/usr/bin/env python3
"""Populate an existing Japanese functional-test Excel template from Feature Map.

Plan generation never approves tests. Import accepts only revision-bound, individually
evidenced checks. Workbook results, human-entered data and formula cells are protected.
Requires openpyxl; supply the original test_template.xlsx via --template.
"""
from __future__ import annotations

import argparse
from copy import copy, deepcopy
from datetime import datetime
import json
from pathlib import Path
import re
import sys
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.worksheet.hyperlink import Hyperlink

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "skills" / "feature-map-workflow" / "scripts"))
from human_work_plan import case_checkpoints, normalize_v2

TYPE_SLOT = {
    "normal": 1, "condition": 2, "branch": 2, "boundary": 3,
    "validation": 4, "invalid": 4, "data": 5, "database": 5, "db": 5,
    "state": 6, "error": 7, "recovery": 7, "permission": 8,
    "security": 8, "external": 9, "integration": 9,
    "retry": 10, "duplicate": 10,
}
VIEWPOINTS = [
    ("主要正常系", "高", "業務の正常な完了と重要な副作用"),
    ("条件分岐", "高", "業務条件で結果が変わる経路"),
    ("境界値", "中", "境界の直前・境界・直後"),
    ("入力検証", "中", "必須、形式、不正データの不登録"),
    ("DB・データ整合性", "高", "更新前後のレコード、キー、関連データ"),
    ("状態遷移", "高", "許可／禁止される状態遷移"),
    ("エラー・復旧", "高", "部分更新を残さない復旧"),
    ("権限・閲覧範囲", "中", "許可・拒否と権限の境界"),
    ("外部連携", "中", "要求・応答とDB更新"),
    ("再実行・二重操作", "中", "再実行時の重複防止"),
]
STATUSES = {"passed": "合格", "failed": "不合格",
            "blocked": "保留", "not-run": "未実施"}


def cell_text(value, length=3200):
    value = str(value or "").strip()[:length]
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def xml_text(node, tag):
    if node is None:
        return ""
    return " ".join((node.findtext(tag) or "").split())


def parse_xml(path):
    raw = path.read_bytes()
    if len(raw) > 2 * 1024 * 1024 or re.search(br"<!DOCTYPE\b|<!ENTITY\b", raw, flags=re.I):
        raise ValueError("Unsafe or oversized Feature Map XML")
    root = ET.fromstring(raw)
    if root.tag != "featureMap" or not xml_text(root, "./meta/feature"):
        raise ValueError("Expected featureMap with meta/feature")
    return root


def load_json(path):
    if path.stat().st_size > 256 * 1024:
        raise ValueError("JSON exceeds 256 KiB")
    return json.loads(path.read_text(encoding="utf-8"))


def dest_check(source, target):
    if target.is_symlink() or source.resolve() == target.resolve() or target.exists():
        raise ValueError("Output must be new, distinct from input, and not a symlink")
    if target.suffix.lower() != ".xlsx":
        raise ValueError("Expected .xlsx output")


def workbook(source):
    book = load_workbook(source)
    for name in ("表紙", "検証リスト", "FT-01"):
        if name not in book:
            raise ValueError(f"Template missing {name}")
    if (book["FT-01"]["H30"].data_type != "f"
            or book["FT-01"]["B31"].data_type != "f"):
        raise ValueError("Unsupported workbook template / formulas missing")
    return book


def clone_sheet(book, index):
    name = f"FT-{index:02d}"
    if name in book:
        return book[name]
    sample = book["FT-01"]
    sheet = book.copy_worksheet(sample)
    sheet.title = name
    sheet.data_validations = deepcopy(sample.data_validations)
    sheet.sheet_properties = deepcopy(sample.sheet_properties)
    sheet.page_margins = copy(sample.page_margins)
    sheet.page_setup = copy(sample.page_setup)
    sheet.print_options = copy(sample.print_options)
    sheet.views = deepcopy(sample.views)
    sheet.print_area = "A1:J90"
    sheet.freeze_panes = sample.freeze_panes
    title, priority, hint = VIEWPOINTS[index - 1] if index <= 10 else ("追加検証", "中", "変更条件と重要な副作用")
    sheet["A1"] = f"{name}  /  {title}"
    sheet["B5"] = name
    sheet["D5"] = title
    sheet["D6"] = priority
    sheet["A34"] = f"観点のヒント：{hint}"
    sheet["H31"].hyperlink = Hyperlink(ref="H31", location="'検証リスト'!A1")
    return sheet


def index_row(sheet, index):
    row = index + 6
    name = f"FT-{index:02d}"
    for col, target in {
        "A": "B5", "C": "D5", "D": "B6", "E": "D6", "F": "F6",
        "G": "B31", "H": "H30", "I": "H6", "L": "B28",
    }.items():
        sheet[f"{col}{row}"] = f"='{name}'!{target}"
    sheet[f"B{row}"] = f'=IF(\'{name}\'!F5="","（機能／シナリオ未入力）",\'{name}\'!F5)'
    sheet[f"J{row}"] = f'=IF(\'{name}\'!B30="","",\'{name}\'!B30)'
    sheet[f"K{row}"] = "開く →"
    sheet[f"K{row}"].hyperlink = Hyperlink(ref=f"K{row}", location=f"'{name}'!A1")


def expand(book, count):
    extra = max(0, count - 10)
    listing = book["検証リスト"]
    if extra:
        # openpyxl does not automatically translate merged regions on row insertion.
        moved = [(str(r), r.min_row, r.max_row, r.min_col, r.max_col)
                 for r in listing.merged_cells.ranges if r.min_row >= 17]
        for addr, *_ in moved:
            listing.unmerge_cells(addr)
        listing.insert_rows(17, extra)
        for _, top, bottom, left, right in moved:
            listing.merge_cells(start_row=top + extra, end_row=bottom + extra,
                                start_column=left, end_column=right)
        for index in range(11, count + 1):
            row = index + 6
            for col in range(1, 13):
                source = listing.cell(16, col)
                target = listing.cell(row, col)
                if source.has_style:
                    target._style = copy(source._style)
            listing.row_dimensions[row].height = listing.row_dimensions[16].height
            index_row(listing, index)
        listing.print_area = f"A1:L{23 + extra}"
    for index in range(2, max(10, count) + 1):
        clone_sheet(book, index)


def choose_cases(root, wanted):
    available = root.findall("./verify/case")
    ids = [c.get("id", "") for c in available]
    if not ids or not all(ids) or len(set(ids)) != len(ids):
        raise ValueError("Feature Map requires unique nonempty verify/case IDs")
    all_cases = dict(zip(ids, available))
    selected = wanted or ids
    if len(set(selected)) != len(selected) or any(v not in all_cases for v in selected):
        raise ValueError("Unknown or repeated --case ID")
    slots = {}
    for ident in selected:
        case = all_cases[ident]
        candidate = TYPE_SLOT.get(case.get("type", "").lower())
        slot = candidate if candidate and candidate not in slots else max([10] + list(slots)) + 1
        slots[slot] = case
    return slots


def human_steps(path):
    if not path:
        return {}
    plan = load_json(path)
    if not isinstance(plan, dict) or plan.get("version") != 1:
        raise ValueError("Invalid Human Work Plan v1")
    steps = {}
    for session in plan.get("sessions", []):
        if not isinstance(session, dict):
            raise ValueError("Invalid session")
        for step in session.get("steps", []):
            if not isinstance(step, dict) or not step.get("caseId") or step["caseId"] in steps:
                raise ValueError("Duplicated or missing caseId in Human Work Plan")
            steps[step["caseId"]] = (session.get("sharedSetup", []), step)
    return steps


def plan_cell(value):
    if len(str(value).strip()) > 3200:
        raise ValueError("Procedure checkpoint exceeds workbook cell capacity; split the case")
    return cell_text(value)


def create_plan(args):
    root = parse_xml(args.feature_map)
    cases = choose_cases(root, args.case)
    usecases = root.findall("./useCase")
    if len(usecases) > 1 and (not args.use_case or not args.case):
        raise ValueError("Multi-use-case XML requires explicit --use-case and --case")
    if args.use_case:
        matching = [u for u in usecases if u.get("id") == args.use_case]
        if len(matching) != 1:
            raise ValueError("Unknown --use-case")
        focus = matching[0]
    else:
        focus = usecases[0] if usecases else None
    raw_plan = load_json(args.human_plan) if args.human_plan else None
    procedure_plan = (normalize_v2(raw_plan, [c.get("id") for c in cases.values()], args.revision)
                      if isinstance(raw_plan, dict) and raw_plan.get("version") == 2 else None)
    if procedure_plan and args.environment is not None and args.environment != procedure_plan["environment"]:
        raise ValueError("Workbook environment does not match procedure applicability")
    checkpoints = case_checkpoints(procedure_plan) if procedure_plan else {}
    if any(check["stepNo"] > 6 for entries in checkpoints.values() for _, _, check in entries):
        raise ValueError("Workbook supports six steps per case; split the case without dropping checks")
    work = {} if procedure_plan else human_steps(args.human_plan)
    if set(work) - {c.get("id") for c in cases.values()}:
        raise ValueError("Human Work Plan includes unselected verification IDs")
    dest_check(args.template, args.output)
    book = workbook(args.template)
    expand(book, max([10] + list(cases)))
    cover = book["表紙"]
    cover["D7"] = cell_text(args.project or xml_text(root, "./meta/feature"))
    cover["D9"] = cell_text(xml_text(root, "./meta/system"))
    cover["D11"] = cell_text(xml_text(root, "./goal/purpose"))
    cover["D15"] = cell_text(xml_text(root, "./goal/scope"))
    cover["D19"] = cell_text(procedure_plan["environment"] if procedure_plan else args.environment or "検証環境未確定")
    cover["K19"] = cell_text(args.revision)
    for slot, case in cases.items():
        ws = book[f"FT-{slot:02d}"]
        kind = TYPE_SLOT.get(case.get("type", "").lower())
        if slot > 10 and kind:
            title, priority, hint = VIEWPOINTS[kind - 1]
            ws["D5"], ws["D6"] = title, priority
            ws["A1"] = f"{ws.title}  /  {title}"
            ws["A34"] = f"観点のヒント：{hint}"
        cid = case.get("id")
        ucid = focus.get("id", "") if focus is not None else ""
        expected = xml_text(case, "expect")
        if not expected:
            raise ValueError(f"{cid} has no expected result; cannot invent an oracle")
        setup, step = work.get(cid, ([], {}))
        ws["F5"] = cell_text(xml_text(focus, "scenario") or xml_text(root, "./meta/feature"))
        ws["B6"] = cell_text(ucid)
        ws["B7"] = cell_text(f"Feature Map: {args.feature_map.name} | UC:{ucid} | VERIFY:{cid}")
        ws["F6"] = "未判定"  # A human must accept the test purpose and expected result.
        ws["B8"] = cell_text(" / ".join(setup) if setup else "要確認：環境・権限・テストデータ")
        ws["B9"] = cell_text(xml_text(focus, "input"))
        ws["B10"] = cell_text(step.get("resetBefore") or "要確認：独立した初期状態")
        ws["B11"] = cell_text(xml_text(focus, "outcome"))
        ws["B12"] = cell_text(f"{cid}：仕様上の期待結果を人が確認してから採用する")
        ws["B16"] = cell_text(step.get("action") or xml_text(case, "condition"))
        ws["C16"] = cell_text(xml_text(focus, "input"))
        if re.search(r"DB|テーブル|登録|更新|行|INSERT|UPDATE|DELETE", expected, re.I):
            ws["E16"] = cell_text(expected)
        else:
            ws["D16"] = cell_text(expected)
        if procedure_plan:
            entries = checkpoints[cid]
            ws["B8"] = plan_cell(" / ".join(dict.fromkeys(
                setup for session, _, _ in entries for setup in session["sharedSetup"])))
            ws["B10"] = plan_cell(" / ".join(f'{step["id"]}: {step["procedure"]["resetBefore"]}'
                                             for _, step, _ in entries))
            # The v1 case row is replaced by individually mapped planned checkpoints.
            for address in ("B16", "C16", "D16", "E16"):
                ws[address] = None
            for _, operation, check in entries:
                row = check["stepNo"] + 15
                part = operation["procedure"]
                ws[f"B{row}"] = plan_cell(
                    f'CONDITION:{check["conditionId"]} | OPERATION:{operation["id"]} | '
                    f'PROCEDURE:{part["id"]}@{part["revision"]}\n{part["action"]}\n'
                    f'要件: {check["requirementRef"]} / 手順正本: {part["sourceRef"]}\n'
                    f'前提: {part["precondition"]} / 初期化: {part["resetBefore"]}\n'
                    f'終了状態: {part["postcondition"]} / 復旧: {part["recovery"]}\n'
                    f'依存: {", ".join(operation["dependsOn"])} / 再開: {operation["restartFrom"]}\n'
                    f'試行: {operation["trial"]["mode"]} / 証跡取得: {check["capture"]}')
                ws[f"C{row}"] = plan_cell(operation["data"])
                column = "E" if re.search(r"DB|テーブル|登録|更新|行|INSERT|UPDATE|DELETE", check["expected"], re.I) else "D"
                ws[f"{column}{row}"] = plan_cell(check["expected"])
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    args.output.parent.mkdir(parents=True, exist_ok=True)
    book.save(args.output)
    print(f"Created {args.output}: {len(cases)} cases. Status remains unapproved/unexecuted.")


def case_sheet(book, case_id):
    matches = []
    for ws in book:
        if not re.fullmatch(r"FT-\d+", ws.title):
            continue
        found = re.search(r"(?:^|\|)\s*VERIFY:([^|]+)", str(ws["B7"].value or ""))
        if found and found.group(1).strip() == case_id:
            matches.append(ws)
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one Excel sheet for verify ID {case_id}")
    return matches[0]


def evidence(item, result_path):
    path = item if isinstance(item, str) else item.get("path") if isinstance(item, dict) else None
    if not isinstance(path, str) or not path.strip():
        raise ValueError("Evidence requires a relative path")
    rel = path.strip().replace("\\", "/")
    if (rel.startswith(("/", "~")) or re.match(r"^[A-Za-z]+:", rel)
            or ".." in rel.split("/")):
        raise ValueError("Unsafe evidence path")
    file = result_path.parent / rel
    if file.is_symlink() or not file.resolve().is_relative_to(result_path.parent.resolve()) or not file.is_file():
        raise ValueError(f"Missing or unsafe evidence file: {rel}")
    label = item.get("description", "") if isinstance(item, dict) else ""
    kind = item.get("type") if isinstance(item, dict) else None
    if kind not in ("画像", "ログ", "SQL", "その他"):
        kind = ("画像" if rel.lower().endswith((".png", ".jpg", ".jpeg"))
                else "SQL" if rel.lower().endswith(".sql") else "ログ")
    return rel, kind, cell_text(label or rel)


def import_one(ws, record, path):
    if (record.get("status") not in STATUSES or not record.get("runId")
            or not isinstance(record.get("checks", []), list)):
        raise ValueError("Result requires status, runId and checks[]")
    if ws["B26"].value:
        raise ValueError("Existing run summary: no overwriting manual evidence")
    checks = {}
    for check in record.get("checks", []):
        if not isinstance(check, dict):
            raise ValueError("Invalid check")
        step_no = check.get("stepNo")
        if step_no is None:
            continue
        if (type(step_no) is not int or step_no not in range(1, 7)
                or step_no in checks or check.get("status") not in STATUSES
                or not isinstance(check.get("evidence"), list) or not check["evidence"]):
            raise ValueError("Each mapped check must identify a unique step, status and evidence")
        if not ws[f"B{step_no + 15}"].value:
            raise ValueError("Cannot import a result for an unplanned step")
        marker = re.match(r"^CONDITION:([^\s|]+) \| OPERATION:", str(ws[f"B{step_no + 15}"].value))
        if marker and check.get("conditionId") != marker.group(1):
            raise ValueError("Result conditionId does not match the planned checkpoint")
        checks[step_no] = check
    if record["status"] == "not-run" and checks:
        raise ValueError("Not-run result cannot include executed steps")
    if record["status"] == "passed" and any(c["status"] != "passed" for c in checks.values()):
        raise ValueError("Run PASS conflicts with a failing/blocked step")
    all_evidence = record.get("evidence", [])
    if not isinstance(all_evidence, list) or len(all_evidence) > 10:
        raise ValueError("At most 10 evidence references per sheet")
    ids = {}
    for index, item in enumerate(all_evidence):
        ref, kind, description = evidence(item, path)
        if ref in ids:
            raise ValueError("Duplicate evidence reference")
        row = index + 40
        if ws[f"F{row}"].value:
            raise ValueError("Existing evidence must not be overwritten")
        eid = f"E-{index + 1:02d}"
        ids[ref] = (eid, row)
        ws[f"C{row}"], ws[f"D{row}"], ws[f"F{row}"] = kind, description, cell_text(ref)
    for number, check in checks.items():
        row = number + 15
        if any(ws[f"{col}{row}"].value for col in ("F", "G", "H", "I")):
            raise ValueError("Refusing to overwrite human-entered step results")
        evidence_ids = []
        for ref in check["evidence"]:
            if not isinstance(ref, str) or ref not in ids:
                raise ValueError("Check evidence is missing from run evidence list")
            eid, evid_row = ids[ref]
            if ws[f"B{evid_row}"].value not in (None, number):
                raise ValueError("Evidence shared by different steps without an independent record")
            ws[f"B{evid_row}"] = number
            evidence_ids.append(eid)
        ws[f"F{row}"] = cell_text(check.get("observedUI", "")) or None
        ws[f"G{row}"] = cell_text(check.get("observedEffect", "")) or None
        ws[f"H{row}"] = STATUSES[check["status"]]
        ws[f"I{row}"] = ", ".join(evidence_ids)
    ws["B26"] = cell_text(f"Run {record['runId']} [{record['status']}] / {path.name}"
                          "（個別証跡のある手順だけ反映）")
    if record.get("timestamp"):
        ws["B30"] = datetime.fromisoformat(record["timestamp"].replace("Z", "+00:00")).date()
        ws["B30"].number_format = "yyyy/mm/dd"
    ws["D30"] = cell_text(record.get("environment"))
    if record.get("operator"):
        ws["H6"] = cell_text(record["operator"])


def import_results(args):
    dest_check(args.workbook, args.output)
    book = workbook(args.workbook)
    revision = str(book["表紙"]["K19"].value or "").strip()
    if not revision:
        raise ValueError("No source revision in workbook 表紙!K19")
    seen = set()
    for path in args.result:
        rec = load_json(path)
        if not isinstance(rec, dict) or rec.get("schemaVersion") != 1:
            raise ValueError("Expected result.json schemaVersion=1")
        cid = rec.get("verifyCaseId")
        if not isinstance(cid, str) or not cid or cid in seen:
            raise ValueError("Missing or repeated verifyCaseId")
        seen.add(cid)
        if rec.get("sourceRevision") != revision:
            raise ValueError("Result sourceRevision does not match the workbook")
        import_one(case_sheet(book, cid), rec, path)
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    args.output.parent.mkdir(parents=True, exist_ok=True)
    book.save(args.output)
    print(f"Imported {len(seen)} result run(s) into {args.output}")


def main():
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest="command", required=True)
    a = sub.add_parser("plan")
    a.add_argument("--template", required=True, type=Path, help="Original test_template.xlsx")
    a.add_argument("--feature-map", required=True, type=Path)
    a.add_argument("--revision", required=True, help="Exact tested source revision")
    a.add_argument("--output", required=True, type=Path)
    a.add_argument("--project", default="")
    a.add_argument("--environment")
    a.add_argument("--human-plan", type=Path)
    a.add_argument("--use-case")
    a.add_argument("--case", action="append", default=[])
    b = sub.add_parser("results")
    b.add_argument("--workbook", required=True, type=Path)
    b.add_argument("--result", action="append", required=True, type=Path)
    b.add_argument("--output", required=True, type=Path)
    args = p.parse_args()
    try:
        if args.command == "plan":
            create_plan(args)
        else:
            import_results(args)
    except (OSError, ValueError, TypeError, ET.ParseError) as exc:
        p.error(str(exc))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
